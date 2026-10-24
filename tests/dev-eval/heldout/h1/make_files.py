#!/usr/bin/env python3
"""files/ と inner/ の docx・xlsx を作る（python-docx 1.2.0、openpyxl 3.1.5）。

使い方（スクラッチの仮想環境で）:
    python make_files.py

数も名前も日付も架空である。同じ元のデータから、示すファイル（files/）、
内側の EVAL の入力（inner/inputs/）、正しい成果物の見本（inner/golden/）を作る。
"""

import datetime
import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
FIXED_TIME = datetime.datetime(2026, 9, 30, 9, 0, 0)

ISSUER = ["サンプル食材卸株式会社", "登録番号：T0000000000000", "〒000-0000 サンプル県サンプル市1-2-3", "TEL 000-0000-0000"]
OLD_BANK = "サンプル銀行 本店 普通 1234567 サンプルショクザイオロシ（カ"
NEW_BANK = "テスト信用金庫 駅前支店 普通 7654321 サンプルショクザイオロシ（カ"

# 得意先一覧：コード、名前、敬称、締め日、支払条件
CUSTOMERS = [
    ("C007", "みなと食堂", "御中", "月末", "翌月末"),
    ("C012", "あおば亭", "御中", "月末", "翌月末"),
    ("C018", "麺処かもめ", "御中", "月末", "翌月末"),
    ("C031", "ひだまりベーカリー", "様", "月末", "翌々月10日"),
    ("C044", "もくれん珈琲", "様", "月末", "翌月末"),
    ("C052", "ホテルさくら", "御中", "月末", "翌月末"),
]
NAME_OF = {c[0]: c[1] for c in CUSTOMERS}

# 出荷明細：(出荷日, 得意先コード, 品名, 単位, 数量, 単価, 税率)
SEP = [
    ("2026/09/02", "C012", "鶏もも肉", "kg", 12, 1180, 8),
    ("2026/09/02", "C012", "豚バラ肉", "kg", 8, 1420, 8),
    ("2026/09/09", "C012", "国産米 10kg", "袋", 3, 4380, 8),
    ("2026/09/09", "C012", "割り箸 100膳", "箱", 4, 785, 10),
    ("2026/09/14", "C012", "卵 M玉 10個入", "パック", 20, 238, 8),
    ("2026/09/14", "C012", "使い捨て手袋", "箱", 5, 645, 10),
    ("2026/09/17", "C012", "卵 M玉 10個入（返品）", "パック", -6, 238, 8),
    ("2026/09/23", "C012", "食用油 一斗缶", "缶", 2, 4180, 8),
    ("2026/09/23", "C012", "生ビール樽 10L", "樽", 2, 9800, 10),
    ("2026/09/05", "C031", "強力粉 25kg", "袋", 4, 5480, 8),
    ("2026/09/05", "C031", "バター 450g", "個", 8, 1980, 8),
    ("2026/09/12", "C031", "砂糖 1kg", "袋", 10, 248, 8),
    ("2026/09/12", "C031", "生クリーム 1L", "本", 6, 880, 8),
    ("2026/09/19", "C031", "ペーパータオル", "袋", 3, 865, 10),
    ("2026/09/26", "C031", "卵 M玉 10個入", "パック", 24, 238, 8),
    ("2026/09/03", "C007", "国産米 10kg", "袋", 5, 4380, 8),
    ("2026/09/03", "C007", "鶏もも肉", "kg", 10, 1180, 8),
    ("2026/09/10", "C007", "食用油 一斗缶", "缶", 1, 4180, 8),
    ("2026/09/10", "C007", "割り箸 100膳", "箱", 3, 785, 10),
    ("2026/09/24", "C007", "生ビール樽 10L", "樽", 2, 9800, 10),
    ("2026/09/04", "C044", "牛乳 1L", "本", 30, 198, 8),
    ("2026/09/11", "C044", "砂糖 1kg", "袋", 8, 248, 8),
    ("2026/09/18", "C044", "使い捨て手袋", "箱", 2, 645, 10),
    ("2026/09/25", "C044", "持ち帰り容器 50個", "袋", 5, 1155, 10),
]
OCT = [
    ("2026/10/01", "C012", "鶏もも肉", "kg", 14, 1180, 8),
    ("2026/10/01", "C012", "豚バラ肉", "kg", 9, 1420, 8),
    ("2026/10/08", "C012", "国産米 10kg", "袋", 4, 4380, 8),
    ("2026/10/08", "C012", "割り箸 100膳", "箱", 6, 785, 10),
    ("2026/10/15", "C012", "小麦粉 25kg", "袋", 2, 5150, 8),
    ("2026/10/15", "C012", "持ち帰り容器 50個", "袋", 8, 1155, 10),
    ("2026/10/22", "C012", "トマト缶 2.5kg", "缶", 6, 780, 8),
    ("2026/10/22", "C012", "牛乳 1L", "本", 24, 198, 8),
    ("2026/10/27", "C012", "生ビール樽 10L", "樽", 3, 9800, 10),
    ("2026/10/06", "C031", "強力粉 25kg", "袋", 6, 5480, 8),
    ("2026/10/06", "C031", "バター 450g", "個", 10, 1980, 8),
    ("2026/10/13", "C031", "砂糖 1kg", "袋", 15, 248, 8),
    ("2026/10/13", "C031", "生クリーム 1L", "本", 9, 880, 8),
    ("2026/10/16", "C031", "バター 450g（返品）", "個", -2, 1980, 8),
    ("2026/10/20", "C031", "ペーパータオル", "袋", 4, 865, 10),
    ("2026/10/20", "C031", "持ち帰り容器 50個", "袋", 7, 1155, 10),
    ("2026/10/28", "C031", "卵 M玉 10個入", "パック", 31, 238, 8),
]


def yen(n):
    """金額の書き方。マイナスは、前に△を付ける（「-」は使わない）"""
    return f"△{abs(n):,}" if n < 0 else f"{n:,}"


def totals(rows):
    """税率ごとの小計と消費税（税率ごとの小計にかけて、1円未満は切り捨て）"""
    sub = {8: 0, 10: 0}
    for row in rows:
        sub[row[6]] += row[4] * row[5]
    tax = {r: sub[r] * r // 100 for r in sub}
    return sub, tax


# ---- xlsx ---------------------------------------------------------------------------------------

def make_shipments(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "出荷明細"
    ws.append(["出荷日", "伝票番号", "得意先コード", "得意先名", "品名", "数量", "単位", "単価", "金額", "税率(%)"])
    counter = {}
    for row in sorted(rows, key=lambda r: (r[0], r[1])):
        date, code, item, unit, qty, price, rate = row
        key = date[5:].replace("/", "")
        counter[key] = counter.get(key, 0) + 1
        ws.append([date, f"S{key}-{counter[key]:02d}", code, NAME_OF[code], item, qty, unit, price, qty * price, rate])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate([12, 11, 12, 18, 24, 8, 8, 8, 10, 9], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    for r in ws.iter_rows(min_row=2, min_col=6, max_col=9):
        for cell in r:
            cell.number_format = "#,##0"
    wb.properties.creator = "sample"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.save(path)


def make_customers(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "得意先一覧"
    ws.append(["得意先コード", "得意先名", "敬称", "締め日", "支払条件"])
    for c in CUSTOMERS:
        ws.append(list(c))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, width in enumerate([14, 22, 8, 8, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.properties.creator = "sample"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.save(path)


# ---- docx ---------------------------------------------------------------------------------------

def _font(run, size=None, bold=None):
    run.font.name = "Yu Gothic"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _para(doc, text, size=10.5, bold=False, align=None):
    p = doc.add_paragraph()
    _font(p.add_run(text), size, bold)
    if align:
        p.alignment = align
    return p


def _table(doc, rows, widths_chars=None, header=True, right_cols=()):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            _font(p.add_run(str(text)), 9.5, bold=(header and i == 0))
            if j in right_cols and not (header and i == 0):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return table


def make_invoice(path, *, number, date_text, code, rows, due_text, bank, note=None):
    name = NAME_OF[code]
    honorific = next(c[2] for c in CUSTOMERS if c[0] == code)
    sub, tax = totals(rows)
    pretax = sub[8] + sub[10]
    total = pretax + tax[8] + tax[10]

    doc = Document()
    doc.styles["Normal"].font.name = "Yu Gothic"
    doc.styles["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    _para(doc, "請求書", 20, True, WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, f"請求書番号：{number}")
    _para(doc, f"請求日：{date_text}")
    _para(doc, f"{name} {honorific}", 13, True)
    for line in ISSUER:
        _para(doc, line, 9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, f"ご請求金額：¥{total:,}（税込）", 13, True)
    _para(doc, f"お支払期限：{due_text}")

    detail = [["出荷日", "品名", "数量", "単位", "単価", "金額", "税率"]]
    for date, _, item, unit, qty, price, rate in sorted(rows, key=lambda r: r[0]):
        detail.append([date[5:].replace("/", "/"), item, yen(qty), unit, f"{price:,}", yen(qty * price), f"{rate}%"])
    _table(doc, detail, right_cols=(2, 4, 5))
    _para(doc, "")

    summary = [
        ["8%対象（軽減税率）小計", f"¥{sub[8]:,}"],
        ["消費税（8%）", f"¥{tax[8]:,}"],
        ["10%対象 小計", f"¥{sub[10]:,}"],
        ["消費税（10%）", f"¥{tax[10]:,}"],
        ["税抜小計", f"¥{pretax:,}"],
        ["ご請求金額（税込）", f"¥{total:,}"],
    ]
    _table(doc, summary, header=False, right_cols=(1,))
    _para(doc, "")
    _para(doc, f"振込先：{bank}")
    if note:
        _para(doc, f"備考：{note}")

    cp = doc.core_properties
    cp.author = "sample"
    cp.last_modified_by = "sample"
    cp.created = FIXED_TIME
    cp.modified = FIXED_TIME
    cp.title = ""
    doc.save(path)


def rows_of(rows, code):
    return [r for r in rows if r[1] == code]


def main():
    files = HERE / "files"
    inner = HERE / "inner"
    # 示すファイル（先月＝9月分）
    make_invoice(
        files / "a.docx",
        number="2609-C012",
        date_text="2026年9月30日",
        code="C012",
        rows=rows_of(SEP, "C012"),
        due_text="2026年10月31日（翌月末払い）",
        bank=OLD_BANK,
        note="今月は新商品カタログを同封しています。",
    )
    make_shipments(files / "b.xlsx", SEP)
    make_customers(files / "c.xlsx")
    # 内側の EVAL の入力（今月＝10月分）
    make_shipments(inner / "inputs" / "d.xlsx", OCT)
    make_customers(inner / "inputs" / "e.xlsx")
    # 正しい成果物の見本
    make_invoice(
        inner / "golden" / "out" / "p.docx",
        number="2610-C012",
        date_text="2026年10月31日",
        code="C012",
        rows=rows_of(OCT, "C012"),
        due_text="2026年11月30日",
        bank=NEW_BANK,
    )
    make_invoice(
        inner / "golden" / "out" / "q.docx",
        number="2610-C031",
        date_text="2026年10月31日",
        code="C031",
        rows=rows_of(OCT, "C031"),
        due_text="2026年12月10日",
        bank=NEW_BANK,
    )
    for code in ("C012", "C031"):
        sub, tax = totals(rows_of(OCT, code))
        print(code, "小計", sub, "消費税", tax, "税抜", sub[8] + sub[10])


if __name__ == "__main__":
    main()
