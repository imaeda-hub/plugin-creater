"""議事録の模擬ユーザーが示すファイル（docx）を作る。
使い方：<仮想環境>/bin/python make_files.py [--out 出力先のフォルダ]   （出力先の既定は files/）
部品の版は ../requirements.txt。txt は手で置いてあるので、ここでは作らない。
"""
import argparse
import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

FIXED = datetime.datetime(2026, 10, 12, 9, 0, 0)  # 作るたびに中身が変わらないよう、作成日時を固定する


def make_docx(path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")

    def line(text: str, bold: bool = False) -> None:
        run = doc.add_paragraph().add_run(text)
        run.bold = bold

    line("定例会議 議事録（10/9）", bold=True)
    line("出席者：田中、鈴木、佐藤")
    line("■議題と内容", bold=True)
    line("1. リリース日程")
    line("テストは10/28に終わる見込みです。")
    line("【決定】本番リリースは10/30に決まりました。")
    line("2. 障害対応の振り返り")
    line("ログの保存期間が短く、原因の調査に時間がかかりました。")
    line("【決定】ログの保存期間は90日にすることになりました。")
    line("3. 来月の体制")
    line("新人が1人増えるため、手順書を整えます。各自、内容を確認します。")
    line("■宿題", bold=True)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("担当", "内容", "期限")):
        cell.text = text
    for row in (("田中", "手順書を更新する", "10/25"), ("鈴木", "ログ保存の設定変更を確認する", "期限未定")):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text

    props = doc.core_properties
    props.author = "佐藤"
    props.last_modified_by = "佐藤"
    props.title = "議事録"
    props.created = FIXED
    props.modified = FIXED
    doc.save(path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(Path(__file__).resolve().parent / "files"))
    out = Path(parser.parse_args().out)
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "先週の議事録.docx")
    print(f"作った：{out}")
