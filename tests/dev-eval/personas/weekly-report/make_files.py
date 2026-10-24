"""週報の模擬ユーザーが示すファイル（docx と xlsx）を作る。
使い方：<仮想環境>/bin/python make_files.py [--out 出力先のフォルダ]   （出力先の既定は files/）
部品の版は ../requirements.txt。csv は手で置いてあるので、ここでは作らない。
"""
import argparse
import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Font

FIXED = datetime.datetime(2026, 10, 12, 9, 0, 0)  # 作るたびに中身が変わらないよう、作成日時を固定する


def make_docx(path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")

    def line(text: str, bold: bool = False) -> None:
        run = doc.add_paragraph().add_run(text)
        run.bold = bold

    line("営業チーム 週報（10/5〜10/9）", bold=True)
    line("■今週の売上", bold=True)
    line("・佐藤：31万円")
    line("・鈴木：18万円")
    line("・田中：1万円未満")
    line("合計：50万円")
    line("■来週の予定", bold=True)
    line("・新規のお客様2件を訪問する")
    line("・見積もりを3件出す")
    line("以上")

    props = doc.core_properties
    props.author = "営業チーム"
    props.last_modified_by = "営業チーム"
    props.title = "週報"
    props.created = FIXED
    props.modified = FIXED
    doc.save(path)


def make_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "集計"
    ws.append(["日付", "担当", "金額（円）", "万円"])
    rows = [("2026-10-05", "佐藤", 312000), ("2026-10-06", "鈴木", 187000), ("2026-10-08", "田中", 6000)]
    for i, (day, who, amount) in enumerate(rows, start=2):
        ws.append([day, who, amount, f"=ROUNDDOWN(C{i}/10000,0)"])
    ws.append(["合計", None, "=SUM(C2:C4)", "=ROUNDDOWN(C5/10000,0)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col, width in zip("ABCD", (14, 10, 14, 8)):
        ws.column_dimensions[col].width = width
    wb.properties.creator = "営業チーム"
    wb.properties.lastModifiedBy = "営業チーム"
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.save(path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default=str(Path(__file__).resolve().parent / "files"))
    out = Path(parser.parse_args().out)
    out.mkdir(parents=True, exist_ok=True)
    make_docx(out / "先週の週報.docx")
    make_xlsx(out / "売上集計表.xlsx")
    print(f"作った：{out}")
